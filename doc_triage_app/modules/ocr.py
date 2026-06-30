from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Callable, List, Optional, Tuple

from PIL import Image

from .utils import clean_text


@dataclass
class OCRResult:
    text: str
    method: str
    warnings: List[str]
    pages_processed: int = 0


def _read_text_file(file_bytes: bytes) -> str:
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            return file_bytes.decode(encoding)
        except Exception:
            continue
    return ""


def _read_docx(file_bytes: bytes) -> str:
    try:
        import docx
    except Exception:
        return ""
    try:
        document = docx.Document(BytesIO(file_bytes))
        parts = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts)
    except Exception:
        return ""


def _pdf_digital_text(file_bytes: bytes, max_pages: int) -> Tuple[str, int]:
    try:
        import fitz  # PyMuPDF
    except Exception:
        return "", 0

    try:
        doc = fitz.open(stream=file_bytes, filetype="pdf")
        parts = []
        pages = min(len(doc), max_pages)
        for idx in range(pages):
            text = doc[idx].get_text("text") or ""
            if text.strip():
                parts.append(text)
        return clean_text("\n\n".join(parts)), pages
    except Exception:
        return "", 0


def _pdf_pages_as_images(file_bytes: bytes, max_pages: int, zoom: float = 2.0) -> List[Image.Image]:
    try:
        import fitz
    except Exception:
        return []

    images: List[Image.Image] = []
    try:
        doc = fitz.open(stream=file_bytes, filetype="pdf")
        pages = min(len(doc), max_pages)
        matrix = fitz.Matrix(zoom, zoom)
        for idx in range(pages):
            pix = doc[idx].get_pixmap(matrix=matrix, alpha=False)
            img = Image.open(BytesIO(pix.tobytes("png"))).convert("RGB")
            images.append(img)
    except Exception:
        return []
    return images


def _local_ocr_image(image: Image.Image) -> Tuple[str, Optional[str]]:
    try:
        import pytesseract
        text = pytesseract.image_to_string(image)
        return text, None
    except Exception as exc:
        return "", f"Local Tesseract OCR failed: {exc}"


def _pil_to_png_bytes(image: Image.Image) -> bytes:
    buffer = BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


def extract_text_from_file(
    filename: str,
    file_bytes: bytes,
    mode: str = "Auto",
    ai_ocr_func: Optional[Callable[[bytes], str]] = None,
    max_pages: int = 5,
) -> OCRResult:
    """Extract text from PDFs, images, text files, and docx files.

    mode:
      - Auto: digital text -> local OCR -> AI OCR if available
      - Local OCR only
      - AI OCR only
    """
    ext = Path(filename).suffix.lower()
    warnings: List[str] = []
    mode = mode or "Auto"

    if ext in {".txt", ".csv"}:
        text = clean_text(_read_text_file(file_bytes))
        return OCRResult(text=text, method="Plain text extraction", warnings=warnings, pages_processed=1)

    if ext == ".docx":
        text = clean_text(_read_docx(file_bytes))
        if text:
            return OCRResult(text=text, method="DOCX text extraction", warnings=warnings, pages_processed=1)
        warnings.append("Could not extract DOCX text.")
        return OCRResult(text="", method="DOCX extraction failed", warnings=warnings)

    if ext == ".pdf":
        if mode != "AI OCR only":
            digital_text, pages = _pdf_digital_text(file_bytes, max_pages=max_pages)
            if len(digital_text) >= 80:
                return OCRResult(
                    text=digital_text,
                    method="PDF embedded text extraction",
                    warnings=warnings,
                    pages_processed=pages,
                )

        images = _pdf_pages_as_images(file_bytes, max_pages=max_pages)
        if not images:
            warnings.append("Could not render PDF pages as images.")
            return OCRResult(text="", method="PDF render failed", warnings=warnings)

        parts: List[str] = []

        if mode != "AI OCR only":
            local_errors = []
            for img in images:
                page_text, err = _local_ocr_image(img)
                if err:
                    local_errors.append(err)
                if page_text.strip():
                    parts.append(page_text)
            local_text = clean_text("\n\n".join(parts))
            if len(local_text) >= 60 or mode == "Local OCR only":
                warnings.extend(list(dict.fromkeys(local_errors)))
                return OCRResult(
                    text=local_text,
                    method="Local Tesseract OCR on rendered PDF pages",
                    warnings=warnings,
                    pages_processed=len(images),
                )
            warnings.extend(list(dict.fromkeys(local_errors)))
            warnings.append("Local OCR returned little text; trying AI OCR if configured.")

        if ai_ocr_func is not None and mode != "Local OCR only":
            ai_parts = []
            for img in images:
                try:
                    ai_text = ai_ocr_func(_pil_to_png_bytes(img))
                    if ai_text.strip():
                        ai_parts.append(ai_text)
                except Exception as exc:
                    warnings.append(f"AI OCR failed on a page: {exc}")
            return OCRResult(
                text=clean_text("\n\n".join(ai_parts)),
                method="AI OCR on rendered PDF pages",
                warnings=warnings,
                pages_processed=len(images),
            )

        return OCRResult(
            text=clean_text("\n\n".join(parts)),
            method="OCR attempted",
            warnings=warnings,
            pages_processed=len(images),
        )

    if ext in {".png", ".jpg", ".jpeg", ".webp", ".tif", ".tiff", ".bmp"}:
        try:
            image = Image.open(BytesIO(file_bytes)).convert("RGB")
        except Exception as exc:
            return OCRResult(text="", method="Image open failed", warnings=[f"Could not open image: {exc}"])

        if mode != "AI OCR only":
            local_text, err = _local_ocr_image(image)
            if err:
                warnings.append(err)
            local_text = clean_text(local_text)
            if len(local_text) >= 40 or mode == "Local OCR only":
                return OCRResult(text=local_text, method="Local Tesseract OCR", warnings=warnings, pages_processed=1)
            warnings.append("Local OCR returned little text; trying AI OCR if configured.")

        if ai_ocr_func is not None and mode != "Local OCR only":
            try:
                ai_text = ai_ocr_func(file_bytes)
                return OCRResult(text=clean_text(ai_text), method="AI OCR", warnings=warnings, pages_processed=1)
            except Exception as exc:
                warnings.append(f"AI OCR failed: {exc}")

        return OCRResult(text="", method="Image OCR attempted", warnings=warnings, pages_processed=1)

    return OCRResult(
        text="",
        method="Unsupported file type",
        warnings=[f"Unsupported file extension: {ext}. Try PDF, image, TXT, CSV, or DOCX."],
    )
